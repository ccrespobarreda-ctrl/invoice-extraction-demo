"""Extract structured data from invoice PDFs, validate it, export JSON + Excel."""
import json, re, sys
from datetime import date
from pathlib import Path
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

num = lambda s: float(s.replace(",", ""))

def parse(pdf_path):
    text = pdfplumber.open(pdf_path).pages[0].extract_text()
    lines = text.splitlines()
    g = lambda pat: (m.group(1).strip() if (m := re.search(pat, text, re.M)) else None)
    items = []
    for ln in lines:
        m = re.match(r"^(.+?) (\d+) ([\d,]+\.\d{2}) ([\d,]+\.\d{2})$", ln)
        if m:
            items.append(dict(description=m[1], quantity=int(m[2]), unit_price=num(m[3]), amount=num(m[4])))
    vat_ids = re.findall(r"VAT ID: (\S+)", text)
    po = g(r"PO number ?(.*)$")
    return dict(
        file=Path(pdf_path).name,
        invoice_number=g(r"Invoice no\. (\S+)"),
        vendor_name=lines[0].replace(" INVOICE", "").strip(),
        vendor_vat_id=vat_ids[0] if vat_ids else None,
        buyer_name=g(r"^(.+?) Bank transfer"),
        buyer_vat_id=vat_ids[1] if len(vat_ids) > 1 else None,
        issue_date=g(r"Issue date (\S+)"),
        due_date=g(r"Due date (\S+)"),
        po_number=po or None,
        currency=g(r"Total \((\w{3})\)"),
        line_items=items,
        subtotal=num(g(r"Subtotal ([\d,.]+)")),
        vat_rate=float(g(r"VAT (\d+)%")) / 100,
        vat_amount=num(g(r"VAT \d+% ([\d,.]+)")),
        total=num(g(r"Total \(\w{3}\) ([\d,.]+)")),
    )

def validate(inv):
    issues = []
    lines_sum = round(sum(i["amount"] for i in inv["line_items"]), 2)
    if abs(lines_sum - inv["subtotal"]) > 0.01:
        issues.append(f"Line items sum {lines_sum:,.2f} != subtotal {inv['subtotal']:,.2f}")
    exp_vat = round(inv["subtotal"] * inv["vat_rate"], 2)
    if abs(exp_vat - inv["vat_amount"]) > 0.01:
        issues.append(f"VAT should be {exp_vat:,.2f}, invoice says {inv['vat_amount']:,.2f}")
    exp_total = round(inv["subtotal"] + inv["vat_amount"], 2)
    if abs(exp_total - inv["total"]) > 0.01:
        issues.append(f"Total mismatch: subtotal + VAT = {exp_total:,.2f}, invoice says {inv['total']:,.2f} (diff {inv['total']-exp_total:+,.2f})")
    if inv["due_date"] and inv["issue_date"] and date.fromisoformat(inv["due_date"]) < date.fromisoformat(inv["issue_date"]):
        issues.append(f"Due date {inv['due_date']} is before issue date {inv['issue_date']}")
    if not inv["po_number"]:
        issues.append("Missing PO number")
    inv["status"] = "NEEDS_REVIEW" if issues else "OK"
    inv["review_reasons"] = issues
    return inv

def to_excel(invs, path):
    wb = Workbook(); ws = wb.active; ws.title = "Invoices"
    head = ["Status", "Invoice #", "Vendor", "Vendor VAT", "Issue date", "Due date", "PO", "Currency", "Subtotal", "VAT", "Total", "Review reasons"]
    ws.append(head)
    thin = Side(style="thin", color="D0D5DA")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(vertical="center")
    for inv in invs:
        ws.append([inv["status"], inv["invoice_number"], inv["vendor_name"], inv["vendor_vat_id"], inv["issue_date"], inv["due_date"],
                   inv["po_number"] or "—", inv["currency"], inv["subtotal"], inv["vat_amount"], inv["total"], "\n".join(inv["review_reasons"]) or "—"])
        r = ws.max_row
        ok = inv["status"] == "OK"
        ws.cell(r, 1).fill = PatternFill("solid", fgColor="D1FAE5" if ok else "FEE2E2")
        ws.cell(r, 1).font = Font(bold=True, color="065F46" if ok else "991B1B")
        for c in ws[r]:
            c.border = Border(bottom=thin); c.alignment = Alignment(vertical="top", wrap_text=True)
        for col in (9, 10, 11): ws.cell(r, col).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJKL", [15, 15, 30, 17, 12, 12, 11, 9, 11, 10, 11, 58]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    li = wb.create_sheet("Line items")
    li.append(["Invoice #", "Description", "Qty", "Unit price", "Amount"])
    for c in li[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F2937")
    for inv in invs:
        for i in inv["line_items"]:
            li.append([inv["invoice_number"], i["description"], i["quantity"], i["unit_price"], i["amount"]])
    for col, w in zip("ABCDE", [16, 36, 6, 12, 12]): li.column_dimensions[col].width = w
    wb.save(path)

if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pdfs", nargs="+", type=Path)
    ap.add_argument("--out", type=Path, default=Path("output"))
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    invs = [validate(parse(p)) for p in sorted(args.pdfs)]
    (args.out / "invoices.json").write_text(json.dumps(invs, indent=2, ensure_ascii=False) + "\n")
    to_excel(invs, args.out / "invoices.xlsx")
    for i in invs: print(i["invoice_number"], i["status"], i["review_reasons"])
